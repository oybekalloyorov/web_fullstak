#!/usr/bin/env python3
"""
Taqdimot (slayd) generatori.

  src/8-sinf/taqdimot/NN-nom.txt  ->  8-sinf/taqdimot/NN-nom.html   (brauzerda ochiladigan slaydlar)
                                  ->  8-sinf/taqdimot/pptx/NN-nom.pptx (PowerPoint, python-pptx bo'lsa)

Taqdimot nomi (NN-nom) mavzu fayli bilan bir xil bo'ladi; sarlavha, bob, belgi va tavsif
src/8-sinf/NN-nom.html dagi meta-blokdan olinadi. Birinchi va oxirgi slaydlar avtomatik yasaladi.

Manba formati (slaydlar orasida alohida qatorda  ---  ):

  # Sarlavha               oddiy slayd sarlavhasi
  ## Bo'lim nomi           bo'lim (ajratuvchi) slayd; keyingi qatorlar — kichik izoh
  - matn                   ro'yxat elementi   ("  - matn" — ichki element)
  1. matn                  raqamlangan qadam
  > matn                   o'xshatish (💡) qutisi
  ! matn                   muhim eslatma qutisi
  = matn                   katta markazlashgan ibora
  | a | b |                jadval (birinchi qator — sarlavha)
  ```lang Fayl nomi        kod bloki ... ```
  ::                       ikkinchi ustunni boshlash
  ? savol / o variant / + to'g'ri variant      — test slaydi
  note: matn               o'qituvchi uchun izoh (slaydda ko'rinmaydi)
  boshqa qator             oddiy matn (abzats)

  Qator ichida: **qalin**, `kod`.

Ishga tushirish:
  python3 slides.py           # HTML + PPTX (PPTX uchun: pip install python-pptx)
  python3 slides.py --html    # faqat HTML (build.py ham shuni chaqiradi)
"""
import html
import json
import pathlib
import re
import sys

ROOT = pathlib.Path(__file__).parent
COURSE_DIR = "8-sinf"
SRC = ROOT / "src" / COURSE_DIR / "taqdimot"
OUT = ROOT / COURSE_DIR / "taqdimot"
COURSE_NAME = "8-sinf · Web Full-stack dasturlash"

BOB_COLORS = {"I": "E2552B", "II": "C99400", "III": "2F6FAD", "IV": "1F7A4D", "V": "8E3FBF"}

META_RE = re.compile(r"^\s*<!--(.*?)-->", re.S)


# ---------------------------------------------------------------- tahlil (parse)

def lesson_meta(slug):
    path = ROOT / "src" / COURSE_DIR / f"{slug}.html"
    meta = {}
    m = META_RE.match(path.read_text(encoding="utf-8")) if path.exists() else None
    if m:
        for part in m.group(1).split("|"):
            if ":" in part:
                k, v = part.split(":", 1)
                meta[k.strip()] = v.strip()
    return meta


def bob_key(bob):
    m = re.match(r"\s*([IVX]+)\s+bob", bob or "")
    return m.group(1) if m else "I"


def parse_deck(text):
    slides = []
    for chunk in re.split(r"^---\s*$", text, flags=re.M):
        if chunk.strip():
            slides.append(parse_slide(chunk.strip("\n")))
    return slides


def parse_slide(chunk):
    slide = {"kind": "normal", "title": "", "cols": [[]], "notes": [], "quiz": None}
    lines = chunk.split("\n")
    i = 0
    col = slide["cols"][0]
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        s = line.strip()
        i += 1
        if not s:
            continue
        if s.startswith("```"):
            head = s[3:].strip().split(None, 1)
            lang = head[0] if head else "text"
            ctitle = head[1] if len(head) > 1 else ""
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            i += 1
            col.append({"t": "code", "lang": lang, "title": ctitle, "text": "\n".join(code)})
        elif s.startswith("## ") and not slide["title"]:
            slide["kind"], slide["title"] = "section", s[3:].strip()
        elif s.startswith("# ") and not slide["title"]:
            slide["title"] = s[2:].strip()
        elif s == "::":
            col = []
            slide["cols"].append(col)
        elif s.startswith("note:"):
            slide["notes"].append(s[5:].strip())
        elif s.startswith("? "):
            slide["quiz"] = {"q": s[2:].strip(), "opts": [], "answer": None}
            slide["kind"] = "quiz"
        elif slide["quiz"] is not None and (s.startswith("o ") or s.startswith("+ ")):
            if s.startswith("+ "):
                slide["quiz"]["answer"] = len(slide["quiz"]["opts"])
            slide["quiz"]["opts"].append(s[2:].strip())
        elif s.startswith("- ") or s.startswith("* "):
            level = 1 if (len(line) - len(line.lstrip())) >= 2 else 0
            col.append({"t": "li", "level": level, "text": s[2:].strip()})
        elif re.match(r"\d+\.\s", s):
            n, rest = s.split(".", 1)
            col.append({"t": "step", "n": n, "text": rest.strip()})
        elif s.startswith("> "):
            col.append({"t": "analogy", "text": s[2:].strip()})
        elif s.startswith("! "):
            col.append({"t": "warn", "text": s[2:].strip()})
        elif s.startswith("= "):
            col.append({"t": "big", "text": s[2:].strip()})
        elif s.startswith("|"):
            rows = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            table = []
            for r in rows:
                cells = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", r.strip().strip("|"))]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                table.append(cells)
            col.append({"t": "table", "rows": table})
        else:
            col.append({"t": "p", "text": s})
    slide["cols"] = [c for c in slide["cols"] if c] or [[]]
    return slide


def load_decks():
    decks = []
    for path in sorted(SRC.glob("[0-9]*.txt")):
        slug = path.stem
        meta = lesson_meta(slug)
        decks.append({
            "slug": slug,
            "num": int(slug.split("-", 1)[0]),
            "meta": meta,
            "bob": bob_key(meta.get("bob")),
            "slides": parse_deck(path.read_text(encoding="utf-8")),
        })
    return decks


# ---------------------------------------------------------------- HTML

def inline_html(text):
    out, pos = [], 0
    for m in re.finditer(r"`([^`]+)`|\*\*(.+?)\*\*", text):
        out.append(html.escape(text[pos:m.start()]))
        if m.group(1) is not None:
            out.append(f"<code>{html.escape(m.group(1))}</code>")
        else:
            out.append(f"<b>{inline_html(m.group(2))}</b>")
        pos = m.end()
    out.append(html.escape(text[pos:]))
    return "".join(out)


HL_LANG = {"py": "python", "output": "plaintext", "django": "html", "js": "javascript", "text": "plaintext"}
LANG_LABEL = {"html": "HTML", "css": "CSS", "js": "JavaScript", "javascript": "JavaScript",
              "python": "Python", "py": "Python", "sql": "SQL", "bash": "Terminal", "json": "JSON",
              "output": "Natija", "django": "Django shablon", "text": "Matn", "http": "HTTP", "ini": ".env"}


def block_html(b):
    t = b["t"]
    if t == "code":
        label = b["title"] or LANG_LABEL.get(b["lang"], b["lang"])
        cls = "s-code out" if b["lang"] == "output" else "s-code"
        return (f'<div class="{cls}"><div class="s-code-head">{html.escape(label)}</div>'
                f'<pre><code class="language-{HL_LANG.get(b["lang"], b["lang"])}">{html.escape(b["text"])}</code></pre></div>')
    if t == "table":
        head, *rows = b["rows"]
        h = "".join(f"<th>{inline_html(c)}</th>" for c in head)
        r = "".join("<tr>" + "".join(f"<td>{inline_html(c)}</td>" for c in row) + "</tr>" for row in rows)
        return f'<table class="s-table"><thead><tr>{h}</tr></thead><tbody>{r}</tbody></table>'
    if t == "analogy":
        return f'<div class="s-box analogy"><span class="ico">💡</span><div>{inline_html(b["text"])}</div></div>'
    if t == "warn":
        return f'<div class="s-box warn"><span class="ico">⚠️</span><div>{inline_html(b["text"])}</div></div>'
    if t == "big":
        return f'<div class="s-big">{inline_html(b["text"])}</div>'
    if t == "step":
        return f'<div class="s-step"><span class="n">{html.escape(b["n"])}</span><div>{inline_html(b["text"])}</div></div>'
    if t == "p":
        return f'<p>{inline_html(b["text"])}</p>'
    return ""


def col_html(blocks):
    out, ul = [], []

    def flush():
        if ul:
            out.append('<ul class="s-list">' + "".join(
                f'<li class="lv{b["level"]}">{inline_html(b["text"])}</li>' for b in ul) + "</ul>")
            ul.clear()

    for b in blocks:
        if b["t"] == "li":
            ul.append(b)
        else:
            flush()
            out.append(block_html(b))
    flush()
    return "".join(out)


def slide_html(s, idx):
    notes = html.escape(" ".join(s["notes"]))
    if s["kind"] == "section":
        sub = col_html(s["cols"][0])
        return (f'<section class="slide section" data-notes="{notes}"><div class="sec-inner">'
                f'<div class="sec-label">Bo\'lim</div><h2>{inline_html(s["title"])}</h2>'
                f'<div class="sec-sub">{sub}</div></div></section>')
    body = ""
    if s["kind"] == "quiz":
        q = s["quiz"]
        letters = "ABCDEF"
        opts = "".join(
            f'<button class="opt" data-ok="{1 if j == q["answer"] else 0}"><span class="l">{letters[j]}</span>'
            f'<span>{inline_html(o)}</span></button>' for j, o in enumerate(q["opts"]))
        extra = col_html(s["cols"][0])
        body = (f'<div class="s-body"><div class="col">{extra}<div class="quiz-q">{inline_html(q["q"])}</div>'
                f'<div class="quiz-opts">{opts}</div><div class="quiz-hint">Javobni tanlang — to\'g\'ri variant yashil bo\'ladi</div></div></div>')
        title = s["title"] or "Savol"
    else:
        cols = s["cols"]
        cls = "s-body cols" if len(cols) > 1 else "s-body"
        body = f'<div class="{cls}">' + "".join(f'<div class="col">{col_html(c)}</div>' for c in cols) + "</div>"
        title = s["title"]
    kind = " quiz" if s["kind"] == "quiz" else ""
    return (f'<section class="slide{kind}" data-notes="{notes}"><header class="s-head"><h2>{inline_html(title)}</h2></header>'
            f'{body}</section>')


def deck_html(deck, prev_d, next_d):
    meta = deck["meta"]
    color = BOB_COLORS[deck["bob"]]
    title = meta.get("title", deck["slug"])
    first = (f'<section class="slide cover"><div class="cover-inner"><div class="cover-icon">{meta.get("icon", "📘")}</div>'
             f'<div class="cover-kicker">{deck["num"]}-mavzu · {html.escape(meta.get("bob", ""))}</div>'
             f'<h1>{html.escape(title)}</h1><p>{html.escape(meta.get("desc", ""))}</p>'
             f'<div class="cover-course">{html.escape(COURSE_NAME)}</div></div></section>')
    nxt = (f'<p>Keyingi mavzu: <b>{next_d["num"]}. {html.escape(next_d["meta"].get("title", ""))}</b></p>'
           if next_d else "<p>Kurs yakunlandi — tabriklaymiz! 🎉</p>")
    last = (f'<section class="slide cover end"><div class="cover-inner"><div class="cover-icon">🙌</div>'
            f'<h1>Savollar bormi?</h1>{nxt}'
            f'<div class="cover-course">Batafsil matn, misollar va testlar: <b>{deck["slug"]}.html</b> darsida</div></div></section>')
    slides = [first] + [slide_html(s, i) for i, s in enumerate(deck["slides"])] + [last]
    prev_link = f'{prev_d["slug"]}.html' if prev_d else ""
    next_link = f'{next_d["slug"]}.html' if next_d else ""
    return f"""<!doctype html>
<html lang="uz">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{deck["num"]}-mavzu: {html.escape(title)} — taqdimot</title>
<meta name="description" content="{html.escape(meta.get("desc", ""))}">
<link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><text y='.9em' font-size='90'>🎞️</text></svg>">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;600&display=swap" rel="stylesheet">
<link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/styles/atom-one-dark.min.css">
<link rel="stylesheet" href="../../assets/css/slides.css">
</head>
<body style="--accent:#{color}">
<div class="deck" id="deck">
{chr(10).join(slides)}
</div>
<div class="bar"><div class="bar-fill" id="barFill"></div></div>
<nav class="controls" id="controls">
  <a class="ctl" href="index.html" title="Barcha taqdimotlar">☰</a>
  <a class="ctl" href="../{deck["slug"]}.html" title="Dars matni">📖</a>
  <a class="ctl" href="pptx/{deck["slug"]}.pptx" title="PowerPoint faylini yuklab olish" download>⬇ PPTX</a>
  <button class="ctl" id="prevBtn" title="Oldingi (←)">‹</button>
  <span class="counter" id="counter">1 / 1</span>
  <button class="ctl" id="nextBtn" title="Keyingi (→)">›</button>
  <button class="ctl" id="notesBtn" title="O'qituvchi izohlari (N)">📝</button>
  <button class="ctl" id="fsBtn" title="To'liq ekran (F)">⛶</button>
</nav>
<aside class="notes" id="notes"><b>O'qituvchi uchun izoh:</b> <span id="notesText"></span></aside>
<script>window.DECK_NAV = {json.dumps({"prev": prev_link, "next": next_link})};</script>
<script src="https://cdnjs.cloudflare.com/ajax/libs/highlight.js/11.9.0/highlight.min.js"></script>
<script src="../../assets/js/slides.js"></script>
</body>
</html>
"""


def index_html(decks):
    groups, last = [], None
    for d in decks:
        bob = d["meta"].get("bob", "")
        if bob != last:
            groups.append((bob, d["bob"], []))
            last = bob
        groups[-1][2].append(d)
    parts = []
    for bob, key, ds in groups:
        cards = "".join(
            f'<div class="ix-card"><div class="ix-num">{d["num"]}</div><div class="ix-main">'
            f'<a class="ix-title" href="{d["slug"]}.html">{d["meta"].get("icon", "")} {html.escape(d["meta"].get("title", ""))}</a>'
            f'<div class="ix-meta">{len(d["slides"]) + 2} ta slayd · <a href="{d["slug"]}.html">▶ Ochish</a> · '
            f'<a href="pptx/{d["slug"]}.pptx" download>⬇ PPTX</a> · <a href="../{d["slug"]}.html">📖 Dars matni</a></div></div></div>'
            for d in ds)
        parts.append(f'<h2 style="--accent:#{BOB_COLORS[key]}">{html.escape(bob)}</h2><div class="ix-grid">{cards}</div>')
    total = sum(len(d["slides"]) + 2 for d in decks)
    return f"""<!doctype html>
<html lang="uz">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>8-sinf taqdimotlari</title>
<meta name="description" content="8-sinf Web Full-stack dasturlash kursining har bir mavzusi bo'yicha taqdimotlar (HTML va PowerPoint).">
<link rel="icon" href="data:image/svg+xml,<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 100 100'><text y='.9em' font-size='90'>🎞️</text></svg>">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap" rel="stylesheet">
<link rel="stylesheet" href="../../assets/css/slides.css">
</head>
<body class="ix">
<main class="ix-wrap">
  <a class="ix-back" href="../index.html">← 8-sinf kursiga qaytish</a>
  <h1>🎞️ 8-sinf taqdimotlari</h1>
  <p class="ix-lead">Har bir mavzu uchun alohida taqdimot: <b>{len(decks)} ta taqdimot, {total} ta slayd</b>.
  Brauzerda ochib darsda ko'rsating yoki PowerPoint (.pptx) faylini yuklab oling.</p>
  <div class="ix-help">
    <b>Boshqaruv:</b> <kbd>→</kbd>/<kbd>Probel</kbd> keyingi · <kbd>←</kbd> oldingi · <kbd>F</kbd> to'liq ekran ·
    <kbd>N</kbd> o'qituvchi izohlari · <kbd>Home</kbd>/<kbd>End</kbd> boshi/oxiri · telefonda — suring.
    Test slaydlarida variantni bossangiz, to'g'ri javob ko'rinadi. PDF kerak bo'lsa — <kbd>Ctrl</kbd>+<kbd>P</kbd> → «PDF sifatida saqlash».
  </div>
  {"".join(parts)}
</main>
</body>
</html>
"""


def build_html(decks):
    OUT.mkdir(parents=True, exist_ok=True)
    for i, d in enumerate(decks):
        prev_d = decks[i - 1] if i > 0 else None
        next_d = decks[i + 1] if i + 1 < len(decks) else None
        (OUT / f"{d['slug']}.html").write_text(deck_html(d, prev_d, next_d), encoding="utf-8")
    (OUT / "index.html").write_text(index_html(decks), encoding="utf-8")


# ---------------------------------------------------------------- PPTX

def build_pptx(decks):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    FONT, MONO = "Calibri", "Consolas"
    DARK, MUTED, WHITE = RGBColor(0x1F, 0x23, 0x2B), RGBColor(0x5B, 0x63, 0x70), RGBColor(0xFF, 0xFF, 0xFF)
    CODE_BG, CODE_FG = RGBColor(0x28, 0x2C, 0x34), RGBColor(0xE6, 0xE6, 0xE6)
    W, H = Inches(13.333), Inches(7.5)

    def rgb(hexs):
        return RGBColor.from_string(hexs)

    def tint(hexs, k):
        r, g, b = (int(hexs[i:i + 2], 16) for i in (0, 2, 4))
        return RGBColor(*(int(c + (255 - c) * k) for c in (r, g, b)))

    def add_runs(par, text, size, color=DARK, bold=False, font=FONT):
        pos = 0
        for m in re.finditer(r"`([^`]+)`|\*\*(.+?)\*\*", text):
            if m.start() > pos:
                _run(par, text[pos:m.start()], size, color, bold, font)
            if m.group(1) is not None:
                _run(par, m.group(1), size * 0.92, RGBColor(0xB0, 0x2A, 0x4A), bold, MONO)
            else:
                add_runs(par, m.group(2), size, color, True, font)
            pos = m.end()
        if pos < len(text):
            _run(par, text[pos:], size, color, bold, font)

    def _run(par, text, size, color, bold, font):
        r = par.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color

    def plain(text):
        return re.sub(r"\*\*(.+?)\*\*|`([^`]+)`", lambda m: m.group(1) or m.group(2), text)

    def lines_for(text, size, width_in, mono=False):
        cw = size / 72 * (0.6 if mono else 0.47)  # bitta belgining taxminiy kengligi (dyuym)
        per = max(8, int(width_in / cw))
        n = 0
        for ln in text.split("\n"):
            # `kod` qismlari monospace — kengroq; ularni og'irroq hisoblaymiz
            kod = sum(len(m) for m in re.findall(r"`([^`]+)`", ln))
            uzun = len(plain(ln)) - kod + (kod * 1.25 if not mono else kod)
            n += max(1, -(-int(uzun) // per))
        return n

    def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, Inches(0.05))
        return tb, tf

    def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        sp = slide.shapes.add_shape(shape, x, y, w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
        if line is None:
            sp.line.fill.background()
        else:
            sp.line.color.rgb = line
            sp.line.width = Pt(1.25)
        sp.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = 0.08
        return sp

    # ---- bloklar balandligini o'lchash va chizish
    def measure(b, size, width):
        t = b["t"]
        if t == "li":
            return lines_for(b["text"], size, width - 0.45 - 0.35 * b["level"]) * size * 1.25 / 72 + 0.06
        if t == "p":
            return lines_for(b["text"], size, width) * size * 1.25 / 72 + 0.1
        if t == "step":
            return max(0.5, lines_for(b["text"], size, width - 0.7) * size * 1.25 / 72) + 0.12
        if t in ("analogy", "warn"):
            return lines_for(b["text"], size * 0.95, width - 0.9) * size * 0.95 * 1.25 / 72 + 0.4
        if t == "big":
            return lines_for(b["text"], size * 1.45, width) * size * 1.45 * 1.25 / 72 + 0.3
        if t == "code":
            cs = code_size(b, size, width - 0.4)
            per = max(10, int((width - 0.45) / (cs / 72 * 0.6)))
            n = sum(max(1, -(-len(l) // per)) for l in b["text"].split("\n"))
            return n * cs * 1.32 / 72 + 0.6
        if t == "table":
            ts = size * 0.8
            total = 0
            ncol = len(b["rows"][0])
            for row in b["rows"]:
                total += max(lines_for(c, ts, width / ncol - 0.2) for c in row) * ts * 1.25 / 72 + 0.16
            return total + 0.12
        return 0

    def code_size(b, size, width):
        cs = size * 0.68
        longest = max(len(l) for l in b["text"].split("\n")) or 1
        fit = width / (longest * 0.6) * 72 * 0.97
        return max(min(cs, 12), min(cs, fit))

    def draw(slide, b, x, y, width, size, color):
        t = b["t"]
        h = measure(b, size, width)
        X, Y, Wd = Inches(x), Inches(y), Inches(width)
        if t == "li":
            ind = 0.35 * b["level"]
            _, tf = textbox(slide, Inches(x + 0.45 + ind), Y, Inches(width - 0.45 - ind), Inches(h))
            add_runs(tf.paragraphs[0], b["text"], size if b["level"] == 0 else size * 0.9)
            dot = box(slide, Inches(x + 0.12 + ind), Inches(y + size / 72 * 0.45), Inches(0.13), Inches(0.13),
                      color if b["level"] == 0 else tint(color_hex, 0.45), shape=MSO_SHAPE.OVAL)
        elif t == "p":
            _, tf = textbox(slide, X, Y, Wd, Inches(h))
            add_runs(tf.paragraphs[0], b["text"], size, MUTED if size else DARK)
        elif t == "step":
            c = box(slide, X, Y, Inches(0.5), Inches(0.5), color, shape=MSO_SHAPE.OVAL)
            c.text_frame.margin_left = c.text_frame.margin_right = 0
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _run(p, b["n"], size * 0.8, WHITE, True, FONT)
            c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _, tf = textbox(slide, Inches(x + 0.7), Inches(y + 0.03), Inches(width - 0.7), Inches(h))
            add_runs(tf.paragraphs[0], b["text"], size)
        elif t in ("analogy", "warn"):
            fill = RGBColor(0xEA, 0xF3, 0xFF) if t == "analogy" else RGBColor(0xFF, 0xF4, 0xE0)
            edge = RGBColor(0x3B, 0x82, 0xF6) if t == "analogy" else RGBColor(0xF5, 0x9E, 0x0B)
            box(slide, X, Y, Wd, Inches(h - 0.12), fill)
            box(slide, X, Y, Inches(0.09), Inches(h - 0.12), edge, shape=MSO_SHAPE.RECTANGLE)
            _, tf = textbox(slide, Inches(x + 0.25), Inches(y + 0.1), Inches(width - 0.4), Inches(h - 0.3),
                            MSO_ANCHOR.MIDDLE)
            p = tf.paragraphs[0]
            _run(p, ("💡 " if t == "analogy" else "⚠️ "), size * 0.95, DARK, False, FONT)
            add_runs(p, b["text"], size * 0.95)
        elif t == "big":
            _, tf = textbox(slide, X, Y, Wd, Inches(h), MSO_ANCHOR.MIDDLE)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_runs(p, b["text"], size * 1.45, color, True)
        elif t == "code":
            cs = code_size(b, size, width - 0.4)
            box(slide, X, Y, Wd, Inches(h - 0.12), CODE_BG if b["lang"] != "output" else RGBColor(0xF1, 0xF5, 0xF9),
                None if b["lang"] != "output" else RGBColor(0xCB, 0xD5, 0xE1))
            label = b["title"] or LANG_LABEL.get(b["lang"], b["lang"])
            _, tf = textbox(slide, Inches(x + 0.2), Inches(y + 0.06), Inches(width - 0.4), Inches(0.3))
            _run(tf.paragraphs[0], label, 11, RGBColor(0x9C, 0xA3, 0xAF), True, FONT)
            _, tf = textbox(slide, Inches(x + 0.2), Inches(y + 0.36), Inches(width - 0.4), Inches(h - 0.5))
            fg = CODE_FG if b["lang"] != "output" else DARK
            for k, ln in enumerate(b["text"].split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.line_spacing = 1.0
                code_runs(p, ln or " ", cs, fg, b["lang"])
        elif t == "table":
            rows, ncol = b["rows"], len(b["rows"][0])
            ts = size * 0.8
            gt = slide.shapes.add_table(len(rows), ncol, X, Y, Wd, Inches(h - 0.12)).table
            for r, row in enumerate(rows):
                for c in range(ncol):
                    cell = gt.cell(r, c)
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    cell.margin_top = cell.margin_bottom = Inches(0.04)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = color if r == 0 else (tint(color_hex, 0.92) if r % 2 else WHITE)
                    p = cell.text_frame.paragraphs[0]
                    add_runs(p, row[c] if c < len(row) else "", ts, WHITE if r == 0 else DARK, r == 0)
        return h

    KW = {"python": r"\b(def|class|return|if|elif|else|for|while|in|import|from|as|True|False|None|and|or|not|with|try|except|print|self|lambda|pass|break|continue)\b",
          "js": r"\b(const|let|var|function|return|if|else|for|while|of|in|true|false|null|undefined|new|class|this|document|console|async|await|=>)\b",
          "sql": r"\b(SELECT|FROM|WHERE|INSERT|INTO|VALUES|UPDATE|SET|DELETE|CREATE|TABLE|PRIMARY|KEY|REFERENCES|NOT|NULL|AND|OR|ORDER|BY|GROUP|HAVING|JOIN|LEFT|INNER|ON|AS|LIMIT|DISTINCT|COUNT|SUM|AVG|MIN|MAX|SERIAL|INTEGER|INT|VARCHAR|TEXT|DATE|BOOLEAN|DEFAULT|UNIQUE|ALTER|DROP|INDEX|LIKE|IN|BETWEEN|IS|DESC|ASC|NUMERIC|CHECK)\b"}

    def code_runs(p, ln, size, fg, lang):
        lang = {"py": "python", "javascript": "js", "html": "html"}.get(lang, lang)
        comment = {"python": "#", "bash": "#", "js": "//", "css": "/*", "sql": "--", "html": "<!--", "django": "{#"}.get(lang)
        if comment and ln.lstrip().startswith(comment):
            _run(p, ln, size, RGBColor(0x7F, 0x84, 0x8E), False, MONO)
            return
        pat = KW.get(lang)
        parts = r"""("[^"]*"|'[^']*')""" + (f"|{pat}" if pat else "")
        if lang in ("html", "django"):
            parts = r"""("[^"]*")|(</?[\w-]+|/?>|\{[{%].*?[%}]\})"""
        pos = 0
        for m in re.finditer(parts, ln, flags=0 if lang != "sql" else re.I):
            if m.start() > pos:
                _run(p, ln[pos:m.start()], size, fg, False, MONO)
            tok = m.group(0)
            colr = RGBColor(0x98, 0xC3, 0x79) if tok[:1] in "\"'" else RGBColor(0xC6, 0x78, 0xDD)
            if lang in ("html", "django") and not tok.startswith('"'):
                colr = RGBColor(0xE0, 0x6C, 0x75)
            if fg == DARK:
                colr = DARK
            _run(p, tok, size, colr, False, MONO)
            pos = m.end()
        if pos < len(ln):
            _run(p, ln[pos:], size, fg, False, MONO)

    def frame(slide, title, num, total, color, deck_title):
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = WHITE
        box(slide, 0, 0, W, Inches(0.12), color, shape=MSO_SHAPE.RECTANGLE)
        _, tf = textbox(slide, Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9), MSO_ANCHOR.MIDDLE)
        add_runs(tf.paragraphs[0], title, 30, DARK, True)
        box(slide, Inches(0.62), Inches(1.22), Inches(1.1), Inches(0.06), color, shape=MSO_SHAPE.RECTANGLE)
        _, tf = textbox(slide, Inches(0.6), Inches(7.0), Inches(10), Inches(0.35))
        _run(tf.paragraphs[0], deck_title, 11, MUTED, False, FONT)
        _, tf = textbox(slide, Inches(11.2), Inches(7.0), Inches(1.55), Inches(0.35))
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        _run(tf.paragraphs[0], f"{num} / {total}", 11, MUTED, False, FONT)

    def layout_cols(slide, cols, color, top=1.55, bottom=6.85):
        avail = bottom - top
        n = len(cols)
        gap = 0.4
        width = (12.13 - gap * (n - 1)) / n
        for size in (28, 26, 24, 22, 20, 19, 18, 17, 16, 15, 14, 13, 12):
            heights = [sum(measure(b, size, width) for b in c) + 0.12 * max(0, len(c) - 1) for c in cols]
            if max(heights) <= avail:
                break
        else:
            OVERFLOW.append(f"{current_slug}: {slide_title_now}")
        for k, c in enumerate(cols):
            x = 0.6 + k * (width + gap)
            total = sum(measure(b, size, width) for b in c) + 0.12 * max(0, len(c) - 1)
            y = top + max(0, (avail - total) / 2) * 0.35 if (n == 1 and len(c) <= 3) else top
            for b in c:
                y += draw(slide, b, x, y, width, size, color) + 0.12

    def notes(slide, texts):
        if texts:
            slide.notes_slide.notes_text_frame.text = " ".join(texts)

    def cover(prs, deck, color, end=False, next_d=None):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = color
        box(s, Inches(0.6), Inches(0.6), Inches(12.13), Inches(6.3), WHITE)
        box(s, Inches(0.6), Inches(0.6), Inches(0.18), Inches(6.3), color, shape=MSO_SHAPE.RECTANGLE)
        meta = deck["meta"]
        if not end:
            _, tf = textbox(s, Inches(1.3), Inches(1.1), Inches(11), Inches(1.1))
            _run(tf.paragraphs[0], meta.get("icon", "📘"), 54, DARK, False, "Segoe UI Emoji")
            _, tf = textbox(s, Inches(1.3), Inches(2.3), Inches(11), Inches(0.5))
            _run(tf.paragraphs[0], f'{deck["num"]}-MAVZU  ·  {meta.get("bob", "").upper()}', 15, color, True, FONT)
            _, tf = textbox(s, Inches(1.3), Inches(2.85), Inches(11), Inches(1.9), MSO_ANCHOR.TOP)
            tsize = 44 if len(meta.get("title", "")) < 40 else 36
            add_runs(tf.paragraphs[0], meta.get("title", ""), tsize, DARK, True)
            _, tf = textbox(s, Inches(1.3), Inches(4.75), Inches(10.8), Inches(1.3))
            add_runs(tf.paragraphs[0], meta.get("desc", ""), 18, MUTED)
        else:
            _, tf = textbox(s, Inches(1.3), Inches(1.6), Inches(11), Inches(1.2))
            _run(tf.paragraphs[0], "🙌", 60, DARK, False, "Segoe UI Emoji")
            _, tf = textbox(s, Inches(1.3), Inches(2.9), Inches(11), Inches(1.0))
            _run(tf.paragraphs[0], "Savollar bormi?", 48, DARK, True, FONT)
            _, tf = textbox(s, Inches(1.3), Inches(4.1), Inches(11), Inches(1.2))
            if next_d:
                add_runs(tf.paragraphs[0], f'Keyingi mavzu: **{next_d["num"]}. {next_d["meta"].get("title", "")}**', 22, MUTED)
            else:
                add_runs(tf.paragraphs[0], "Kurs yakunlandi — tabriklaymiz! 🎉", 22, MUTED)
        _, tf = textbox(s, Inches(1.3), Inches(6.2), Inches(11), Inches(0.45))
        _run(tf.paragraphs[0], COURSE_NAME, 14, MUTED, True, FONT)

    def section(prs, sl, color, num, total, deck_title):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = color
        _, tf = textbox(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5))
        _run(tf.paragraphs[0], "BO'LIM", 16, tint(color_hex, 0.7), True, FONT)
        _, tf = textbox(s, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.8), MSO_ANCHOR.TOP)
        add_runs(tf.paragraphs[0], sl["title"], 44, WHITE, True)
        sub = " ".join(b["text"] for b in sl["cols"][0] if "text" in b)
        if sub:
            _, tf = textbox(s, Inches(1.0), Inches(4.3), Inches(11.3), Inches(1.6))
            add_runs(tf.paragraphs[0], sub, 22, tint(color_hex, 0.85))
        _, tf = textbox(s, Inches(11.2), Inches(7.0), Inches(1.55), Inches(0.35))
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
        _run(tf.paragraphs[0], f"{num} / {total}", 11, tint(color_hex, 0.7), False, FONT)
        notes(s, sl["notes"])

    def quiz(prs, sl, color, num, total, deck_title, reveal):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        q = sl["quiz"]
        frame(s, (sl["title"] or "Savol") + (" — javob" if reveal else ""), num, total, color, deck_title)
        y = 1.55
        if sl["cols"][0]:
            pre = sl["cols"][0]
            for b in pre:
                y += draw(s, b, 0.6, y, 12.13, 18, color) + 0.08
        _, tf = textbox(s, Inches(0.6), Inches(y), Inches(12.13), Inches(1.0))
        qsize = 24 if len(q["q"]) < 90 else 20
        add_runs(tf.paragraphs[0], q["q"], qsize, DARK, True)
        y += lines_for(q["q"], qsize, 12.1) * qsize * 1.3 / 72 + 0.25
        n = len(q["opts"])
        two = n == 4 and max(len(plain(o)) for o in q["opts"]) < 46
        oh = min(0.85, (6.75 - y) / ((n + 1) // 2 if two else n) - 0.15)
        osize = 19 if oh > 0.6 else 16
        for j, o in enumerate(q["opts"]):
            if two:
                ox, oy, ow = 0.6 + (j % 2) * 6.17, y + (j // 2) * (oh + 0.15), 5.96
            else:
                ox, oy, ow = 0.6, y + j * (oh + 0.15), 12.13
            ok = reveal and j == q["answer"]
            box(s, Inches(ox), Inches(oy), Inches(ow), Inches(oh),
                RGBColor(0xDC, 0xFC, 0xE7) if ok else RGBColor(0xF6, 0xF7, 0xF9),
                RGBColor(0x16, 0xA3, 0x4A) if ok else RGBColor(0xD9, 0xDE, 0xE5))
            c = box(s, Inches(ox + 0.15), Inches(oy + (oh - 0.46) / 2), Inches(0.46), Inches(0.46),
                    RGBColor(0x16, 0xA3, 0x4A) if ok else color, shape=MSO_SHAPE.OVAL)
            c.text_frame.margin_left = c.text_frame.margin_right = 0
            c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _run(p, "✓" if ok else "ABCDEF"[j], 15, WHITE, True, FONT)
            _, tf = textbox(s, Inches(ox + 0.8), Inches(oy), Inches(ow - 0.95), Inches(oh), MSO_ANCHOR.MIDDLE)
            add_runs(tf.paragraphs[0], o, osize, DARK, ok)
        if not reveal:
            notes(s, [f'To\'g\'ri javob: {"ABCDEF"[q["answer"]]}) {plain(q["opts"][q["answer"]])}.'] + sl["notes"])
        else:
            notes(s, sl["notes"])

    out_dir = OUT / "pptx"
    out_dir.mkdir(parents=True, exist_ok=True)
    for i, deck in enumerate(decks):
        global color_hex
        color_hex = BOB_COLORS[deck["bob"]]
        color = rgb(color_hex)
        prs = Presentation()
        prs.slide_width, prs.slide_height = W, H
        prs.core_properties.title = f'{deck["num"]}-mavzu: {deck["meta"].get("title", "")}'
        prs.core_properties.author = COURSE_NAME
        deck_title = f'{deck["num"]}-mavzu · {deck["meta"].get("title", "")}'
        total = 2 + sum(2 if s["kind"] == "quiz" else 1 for s in deck["slides"])
        num = 1
        cover(prs, deck, color)
        for sl in deck["slides"]:
            num += 1
            if sl["kind"] == "section":
                section(prs, sl, color, num, total, deck_title)
            elif sl["kind"] == "quiz":
                quiz(prs, sl, color, num, total, deck_title, False)
                num += 1
                quiz(prs, sl, color, num, total, deck_title, True)
            else:
                s = prs.slides.add_slide(prs.slide_layouts[6])
                frame(s, sl["title"], num, total, color, deck_title)
                global current_slug, slide_title_now
                current_slug, slide_title_now = deck["slug"], sl["title"]
                layout_cols(s, sl["cols"], color)
                notes(s, sl["notes"])
        cover(prs, deck, color, end=True, next_d=decks[i + 1] if i + 1 < len(decks) else None)
        prs.save(out_dir / f"{deck['slug']}.pptx")


color_hex = "E2552B"
OVERFLOW = []
current_slug = slide_title_now = ""


def build(pptx=True):
    decks = load_decks()
    build_html(decks)
    msg = f"  Taqdimotlar: {len(decks)} ta (HTML)"
    if pptx:
        try:
            build_pptx(decks)
            msg += " + PPTX"
            for o in OVERFLOW:
                print("  ⚠️ sig'madi:", o)
        except ImportError:
            msg += " — PPTX o'tkazib yuborildi (pip install python-pptx)"
    print(msg)
    return decks


if __name__ == "__main__":
    build(pptx="--html" not in sys.argv)
